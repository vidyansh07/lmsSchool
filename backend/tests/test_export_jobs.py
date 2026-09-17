"""Background export jobs — CSV, XLSX and PDF, off the request cycle.

`ReportExportView` (`tests/test_reporting.py`) covers the synchronous CSV
download, which this feature sits beside rather than replaces. What is
specific to a queued job is what is tested here: the state machine
(`QUEUED → PROCESSING → COMPLETED|FAILED|CANCELLED`), that every format
actually opens and is safe to open, that scope is re-derived from the
requester's *current* access rather than trusted from the moment the job was
queued, and that a job belongs to its owner the same way a submission file
does — another user's id 404s rather than 403s.

`CELERY_TASK_ALWAYS_EAGER` is on in `config.settings.test`, so `.delay()`
inside the queue view runs `run_export` synchronously, inline. Most tests
below queue through the API and read the finished job straight back. Tests
that need to inspect an in-between state — a job still `QUEUED`, a scope that
narrows after queueing but before running — patch `tasks.run_export.delay` to
a no-op and call `tasks.run_export(job_id)` directly, the same way
`tests/test_background_work.py` calls `retry_pending_email()` directly rather
than through the queue.
"""

from __future__ import annotations

import csv
import io
from datetime import timedelta
from io import BytesIO

import openpyxl
import pytest
from django.utils import timezone

from apps.accounts.models import UserRole
from apps.audit.models import AuditAction, AuditLog
from apps.reporting import tasks, writers
from apps.reporting.models import ExportFormat, ExportJob, ExportStatus

EXPORTS_URL = "/api/v1/reports/exports/"


def _detail_url(job_id) -> str:
    return f"{EXPORTS_URL}{job_id}/"


def _download_url(job_id) -> str:
    return f"{EXPORTS_URL}{job_id}/download/"


def _cancel_url(job_id) -> str:
    return f"{EXPORTS_URL}{job_id}/cancel/"


def _queue(client, user, *, report_key="student_progress", fmt=ExportFormat.CSV, **extra):
    client.force_login(user)
    payload = {"report_key": report_key, "format": fmt, **extra}
    return client.post(EXPORTS_URL, payload)


def _extra_enrollment(admin_user, batch, *, index: int):
    """A cheap extra enrolment, for tests that need more rows than the shared
    fixtures provide (the PDF row cap, mainly)."""
    from apps.enrollments.services import enrol_student
    from apps.students.services import create_student

    student = create_student(
        email=f"extra{index}@example.test",
        first_name="Extra",
        last_name=f"Student{index}",
        actor=admin_user,
        password="correct-horse-battery-staple",
        send_invitation=False,
    )
    return enrol_student(student=student, batch=batch, actor=admin_user)


@pytest.fixture(autouse=True)
def _forget_policy_memo():
    from apps.common.request_context import clear_scope

    clear_scope()
    yield
    clear_scope()


# ---------------------------------------------------------------------------
# Queue → run → completed, for each format
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_queueing_a_csv_export_runs_it_to_completion(api_client_no_csrf, admin_user, enrollment):
    response = _queue(api_client_no_csrf, admin_user)

    assert response.status_code == 202
    body = response.json()
    assert body["status"] == ExportStatus.COMPLETED

    job = ExportJob.objects.get(pk=body["id"])
    assert job.status == ExportStatus.COMPLETED
    assert job.row_count >= 1
    assert job.checksum
    assert job.size_bytes > 0
    assert job.file


@pytest.mark.django_db
def test_the_completed_csv_file_parses_and_has_a_header(api_client_no_csrf, admin_user, enrollment):
    response = _queue(api_client_no_csrf, admin_user)
    job_id = response.json()["id"]

    download = api_client_no_csrf.get(_download_url(job_id))
    assert download.status_code == 200
    assert download["Content-Type"].startswith("text/csv")

    body = b"".join(download.streaming_content).decode("utf-8")
    rows = list(csv.reader(io.StringIO(body)))
    assert rows[0][0] == "Student"
    assert len(rows) > 1


@pytest.mark.django_db
def test_queueing_an_xlsx_export_produces_a_file_openpyxl_can_read(
    api_client_no_csrf, admin_user, enrollment
):
    response = _queue(api_client_no_csrf, admin_user, fmt=ExportFormat.XLSX)
    job_id = response.json()["id"]

    download = api_client_no_csrf.get(_download_url(job_id))
    assert download.status_code == 200
    assert download["Content-Type"] == writers.CONTENT_TYPES["xlsx"]

    content = b"".join(download.streaming_content)
    workbook = openpyxl.load_workbook(BytesIO(content))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "Student"
    assert len(rows) > 1


@pytest.mark.django_db
def test_queueing_a_pdf_export_produces_a_real_pdf(api_client_no_csrf, admin_user, enrollment):
    response = _queue(api_client_no_csrf, admin_user, fmt=ExportFormat.PDF)
    job_id = response.json()["id"]

    download = api_client_no_csrf.get(_download_url(job_id))
    assert download.status_code == 200
    assert download["Content-Type"] == "application/pdf"

    content = b"".join(download.streaming_content)
    assert content.startswith(b"%PDF")
    assert len(content) > 500


# ---------------------------------------------------------------------------
# Formula injection — the most important test in this file
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_xlsx_neutralises_formula_injection(
    api_client_no_csrf, admin_user, published_course, trainer_profile
):
    """A batch name starting `=cmd|...` must come back as text, not a formula.

    A batch name is only checked for control characters
    (`validate_no_control_characters`), unlike a person's name, so this is a
    value that can legitimately reach the database through the ordinary
    batch-creation path — exactly the kind of value the sanitiser exists to
    catch before it becomes a formula that runs on whoever opens the file.
    """
    from apps.batches.services import create_batch

    today = timezone.localdate()
    malicious = create_batch(
        actor=admin_user,
        name="=cmd|'/c calc'!A1",
        course=published_course,
        trainer=trainer_profile,
        start_date=today - timedelta(days=1),
        end_date=today + timedelta(days=30),
        capacity=5,
    )
    assert malicious.name.startswith("=")

    response = _queue(
        api_client_no_csrf, admin_user, report_key="batch_performance", fmt=ExportFormat.XLSX
    )
    job_id = response.json()["id"]

    download = api_client_no_csrf.get(_download_url(job_id))
    content = b"".join(download.streaming_content)

    workbook = openpyxl.load_workbook(BytesIO(content))
    sheet = workbook.active
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    name_index = header.index("Name")

    found = False
    for row in sheet.iter_rows(min_row=2):
        cell = row[name_index]
        if cell.value and str(cell.value).endswith("cmd|'/c calc'!A1"):
            found = True
            assert cell.data_type == "s", "the malicious cell must be stored as text, not a formula"
            assert str(cell.value).startswith("'="), "sanitise() should have prefixed the '='"
    assert found, "the malicious batch never appeared in the export"


@pytest.mark.django_db
def test_csv_export_of_the_same_malicious_batch_is_also_neutralised(
    api_client_no_csrf, admin_user, published_course, trainer_profile
):
    from apps.batches.services import create_batch

    today = timezone.localdate()
    create_batch(
        actor=admin_user,
        name="=cmd|'/c calc'!A1",
        course=published_course,
        trainer=trainer_profile,
        start_date=today - timedelta(days=1),
        end_date=today + timedelta(days=30),
        capacity=5,
    )
    response = _queue(
        api_client_no_csrf, admin_user, report_key="batch_performance", fmt=ExportFormat.CSV
    )
    job_id = response.json()["id"]
    download = api_client_no_csrf.get(_download_url(job_id))
    body = b"".join(download.streaming_content).decode("utf-8")
    assert "'=cmd|" in body
    assert "\n=cmd|" not in body


# ---------------------------------------------------------------------------
# The PDF row cap
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_the_pdf_row_cap_fails_the_job_without_producing_a_file(
    api_client_no_csrf, admin_user, batch, enrollment, other_enrollment, monkeypatch
):
    monkeypatch.setattr(writers, "PDF_MAX_ROW_LIMIT", 1)

    response = _queue(api_client_no_csrf, admin_user, fmt=ExportFormat.PDF)
    assert response.status_code == 202
    job = ExportJob.objects.get(pk=response.json()["id"])

    assert job.status == ExportStatus.FAILED
    assert not job.file
    assert job.row_count == 0
    assert "PDF" in job.error or "pdf" in job.error.lower()
    assert "Traceback" not in job.error


@pytest.mark.django_db
def test_the_pdf_row_cap_message_says_to_use_another_format(
    api_client_no_csrf, admin_user, enrollment, monkeypatch
):
    monkeypatch.setattr(writers, "PDF_MAX_ROW_LIMIT", 0)

    response = _queue(api_client_no_csrf, admin_user, fmt=ExportFormat.PDF)
    job = ExportJob.objects.get(pk=response.json()["id"])

    assert job.status == ExportStatus.FAILED
    assert job.error
    assert "csv" in job.error.lower() or "excel" in job.error.lower()


@pytest.mark.django_db
def test_a_pdf_export_within_the_cap_still_succeeds(
    api_client_no_csrf, admin_user, enrollment, monkeypatch
):
    monkeypatch.setattr(writers, "PDF_MAX_ROW_LIMIT", 500)
    response = _queue(api_client_no_csrf, admin_user, fmt=ExportFormat.PDF)
    job = ExportJob.objects.get(pk=response.json()["id"])
    assert job.status == ExportStatus.COMPLETED
    assert job.row_count >= 1


# ---------------------------------------------------------------------------
# Queueing is refused up front
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_an_unknown_report_key_is_refused_at_queue_time(api_client_no_csrf, admin_user):
    response = _queue(api_client_no_csrf, admin_user, report_key="not-a-real-report")
    assert response.status_code == 404
    assert not ExportJob.objects.exists()


@pytest.mark.django_db
def test_an_invalid_format_is_rejected(api_client_no_csrf, admin_user):
    response = _queue(api_client_no_csrf, admin_user, fmt="doc")
    assert response.status_code == 400
    assert not ExportJob.objects.exists()


@pytest.mark.django_db
def test_a_batch_id_that_does_not_exist_404s_at_queue_time(api_client_no_csrf, admin_user):
    response = _queue(api_client_no_csrf, admin_user, batch="00000000-0000-0000-0000-000000000000")
    assert response.status_code == 404
    assert not ExportJob.objects.exists()


@pytest.mark.django_db
def test_a_counsellor_may_queue_an_export(api_client_no_csrf, counsellor_user):
    """D-130: holds `report.view_any` like a manager, so the gate opens."""
    response = _queue(api_client_no_csrf, counsellor_user)
    assert response.status_code == 202, response.data
    assert ExportJob.objects.filter(requested_by=counsellor_user).exists()


@pytest.mark.django_db
def test_a_trainer_is_refused_at_queue_time(api_client_no_csrf, trainer_profile, enrollment):
    response = _queue(api_client_no_csrf, trainer_profile.user)
    assert response.status_code == 403
    assert not ExportJob.objects.exists()


@pytest.mark.django_db
def test_an_anonymous_caller_is_refused(api_client_no_csrf):
    response = api_client_no_csrf.post(
        EXPORTS_URL, {"report_key": "student_progress", "format": ExportFormat.CSV}
    )
    assert response.status_code in (401, 403)


@pytest.mark.django_db
def test_a_manager_can_queue_and_download(api_client_no_csrf, manager_user, enrollment):
    response = _queue(api_client_no_csrf, manager_user)
    job_id = response.json()["id"]
    download = api_client_no_csrf.get(_download_url(job_id))
    assert download.status_code == 200


# ---------------------------------------------------------------------------
# Scope is re-derived when the task actually runs, never trusted from queue time
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_scope_is_re_derived_when_the_capability_is_revoked_before_running(
    api_client_no_csrf, admin_user, manager_user, enrollment, monkeypatch
):
    monkeypatch.setattr(tasks.run_export, "delay", lambda *a, **k: None)

    response = _queue(api_client_no_csrf, manager_user)
    job = ExportJob.objects.get(pk=response.json()["id"])
    assert job.status == ExportStatus.QUEUED

    manager_user.role = UserRole.STUDENT
    manager_user.save(update_fields=["role"])

    tasks.run_export(str(job.pk))
    job.refresh_from_db()

    assert job.status == ExportStatus.FAILED
    assert not job.file
    assert job.error
    assert "permission" in job.error.lower()


@pytest.mark.django_db
def test_scope_is_re_derived_when_the_account_is_deactivated_before_running(
    api_client_no_csrf, admin_user, manager_user, enrollment, monkeypatch
):
    monkeypatch.setattr(tasks.run_export, "delay", lambda *a, **k: None)

    response = _queue(api_client_no_csrf, manager_user)
    job = ExportJob.objects.get(pk=response.json()["id"])

    manager_user.is_active = False
    manager_user.save(update_fields=["is_active"])

    tasks.run_export(str(job.pk))
    job.refresh_from_db()

    assert job.status == ExportStatus.FAILED
    assert not job.file
    assert "active" in job.error.lower()


@pytest.mark.django_db
def test_scope_is_re_derived_when_the_scoped_batch_is_no_longer_visible(
    api_client_no_csrf, manager_user, batch, enrollment, monkeypatch
):
    """An export scoped to a batch must fail if that batch stops being visible
    before the worker gets to it — not silently widen to every batch. A
    manager holds `batch.view_any` and would otherwise see every *other* batch
    just fine, which is exactly why this has to check the one the job named,
    not just "does this user see anything at all".
    """
    from apps.common.deletion import soft_delete

    monkeypatch.setattr(tasks.run_export, "delay", lambda *a, **k: None)

    response = _queue(
        api_client_no_csrf, manager_user, report_key="attendance", batch=str(batch.pk)
    )
    assert response.status_code == 202
    job = ExportJob.objects.get(pk=response.json()["id"])
    assert job.filters["batch_id"] == str(batch.pk)

    soft_delete(instance=batch, actor=manager_user, reason="Batch closed.")

    tasks.run_export(str(job.pk))
    job.refresh_from_db()

    assert job.status == ExportStatus.FAILED
    assert not job.file
    assert "batch" in job.error.lower()


@pytest.mark.django_db
def test_run_export_is_a_noop_for_a_job_that_is_not_queued(admin_user, enrollment):
    """`task_acks_late` makes redelivery possible; a second delivery of a task
    for an already-completed job must not run the export twice.
    """
    job = ExportJob.objects.create(
        report_key="student_progress",
        format=ExportFormat.CSV,
        requested_by=admin_user,
        status=ExportStatus.COMPLETED,
    )
    assert tasks.run_export(str(job.pk)) is False
    job.refresh_from_db()
    assert job.status == ExportStatus.COMPLETED
    assert not job.file


@pytest.mark.django_db
def test_run_export_on_a_missing_job_id_does_nothing():
    assert tasks.run_export("00000000-0000-0000-0000-000000000000") is False


# ---------------------------------------------------------------------------
# Ownership: ids that are not yours 404, they do not 403
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_another_users_job_404s_on_detail(api_client_no_csrf, admin_user, manager_user, enrollment):
    response = _queue(api_client_no_csrf, admin_user)
    job_id = response.json()["id"]

    api_client_no_csrf.force_login(manager_user)
    assert api_client_no_csrf.get(_detail_url(job_id)).status_code == 404


@pytest.mark.django_db
def test_another_users_job_404s_on_download(
    api_client_no_csrf, admin_user, manager_user, enrollment
):
    response = _queue(api_client_no_csrf, admin_user)
    job_id = response.json()["id"]

    api_client_no_csrf.force_login(manager_user)
    assert api_client_no_csrf.get(_download_url(job_id)).status_code == 404


@pytest.mark.django_db
def test_another_users_job_404s_on_cancel(api_client_no_csrf, admin_user, manager_user, enrollment):
    response = _queue(api_client_no_csrf, admin_user)
    job_id = response.json()["id"]

    api_client_no_csrf.force_login(manager_user)
    assert api_client_no_csrf.post(_cancel_url(job_id)).status_code == 404


@pytest.mark.django_db
def test_export_view_any_sees_others_jobs_in_the_list(
    api_client_no_csrf, admin_user, manager_user, enrollment
):
    """`admin_user` holds `export.view_any`; `manager_user` does not."""
    response = _queue(api_client_no_csrf, manager_user)
    job_id = response.json()["id"]

    api_client_no_csrf.force_login(admin_user)
    ids = {row["id"] for row in api_client_no_csrf.get(EXPORTS_URL).json()}
    assert job_id in ids


@pytest.mark.django_db
def test_export_view_any_can_view_and_download_someone_elses_job(
    api_client_no_csrf, admin_user, manager_user, enrollment
):
    response = _queue(api_client_no_csrf, manager_user)
    job_id = response.json()["id"]

    api_client_no_csrf.force_login(admin_user)
    assert api_client_no_csrf.get(_detail_url(job_id)).status_code == 200
    assert api_client_no_csrf.get(_download_url(job_id)).status_code == 200


@pytest.mark.django_db
def test_a_caller_without_view_any_sees_only_their_own_jobs_in_the_list(
    api_client_no_csrf, admin_user, manager_user, enrollment
):
    _queue(api_client_no_csrf, admin_user)
    response = _queue(api_client_no_csrf, manager_user)
    own_job_id = response.json()["id"]

    api_client_no_csrf.force_login(manager_user)
    ids = {row["id"] for row in api_client_no_csrf.get(EXPORTS_URL).json()}
    assert ids == {own_job_id}


# ---------------------------------------------------------------------------
# Cancellation
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_cancelling_a_queued_job_produces_no_file(
    api_client_no_csrf, admin_user, enrollment, monkeypatch
):
    monkeypatch.setattr(tasks.run_export, "delay", lambda *a, **k: None)

    response = _queue(api_client_no_csrf, admin_user)
    job_id = response.json()["id"]

    cancel = api_client_no_csrf.post(_cancel_url(job_id))
    assert cancel.status_code == 200
    assert cancel.json()["status"] == ExportStatus.CANCELLED

    job = ExportJob.objects.get(pk=job_id)
    assert job.status == ExportStatus.CANCELLED
    assert not job.file


@pytest.mark.django_db
def test_a_cancelled_job_is_not_run_if_the_task_arrives_late(
    api_client_no_csrf, admin_user, enrollment, monkeypatch
):
    monkeypatch.setattr(tasks.run_export, "delay", lambda *a, **k: None)

    response = _queue(api_client_no_csrf, admin_user)
    job_id = response.json()["id"]
    api_client_no_csrf.post(_cancel_url(job_id))

    assert tasks.run_export(job_id) is False
    job = ExportJob.objects.get(pk=job_id)
    assert job.status == ExportStatus.CANCELLED
    assert not job.file


@pytest.mark.django_db
def test_cancelling_an_already_completed_job_conflicts(api_client_no_csrf, admin_user, enrollment):
    response = _queue(api_client_no_csrf, admin_user)
    job_id = response.json()["id"]
    assert api_client_no_csrf.post(_cancel_url(job_id)).status_code == 409


# ---------------------------------------------------------------------------
# The error field: readable, never a traceback, never leaked to the wrong caller
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_an_unexpected_failure_never_puts_a_traceback_in_the_error_field(
    api_client_no_csrf, admin_user, enrollment, monkeypatch
):
    def _explode(*args, **kwargs):
        raise RuntimeError("token=eyJsecretvalue123 leaked from a stack frame\nline 2\nline 3")

    monkeypatch.setattr(tasks.reports, "run", _explode)

    response = _queue(api_client_no_csrf, admin_user)
    job = ExportJob.objects.get(pk=response.json()["id"])

    assert job.status == ExportStatus.FAILED
    assert job.error == tasks.GENERIC_FAILURE_MESSAGE
    assert "token" not in job.error
    assert "Traceback" not in job.error
    assert "\n" not in job.error
    assert len(job.error) <= 500


# ---------------------------------------------------------------------------
# Download availability
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_download_404s_while_the_job_is_still_queued(
    api_client_no_csrf, admin_user, enrollment, monkeypatch
):
    monkeypatch.setattr(tasks.run_export, "delay", lambda *a, **k: None)
    response = _queue(api_client_no_csrf, admin_user)
    job_id = response.json()["id"]
    assert api_client_no_csrf.get(_download_url(job_id)).status_code == 404


@pytest.mark.django_db
def test_download_404s_for_a_failed_job(api_client_no_csrf, admin_user, enrollment):
    job = ExportJob.objects.create(
        report_key="student_progress",
        format=ExportFormat.CSV,
        requested_by=admin_user,
        status=ExportStatus.FAILED,
        error="Something went wrong.",
    )
    api_client_no_csrf.force_login(admin_user)
    assert api_client_no_csrf.get(_download_url(job.pk)).status_code == 404


@pytest.mark.django_db
def test_download_404s_once_the_job_has_expired(api_client_no_csrf, admin_user, enrollment):
    response = _queue(api_client_no_csrf, admin_user)
    job = ExportJob.objects.get(pk=response.json()["id"])
    job.expires_at = timezone.now() - timedelta(seconds=1)
    job.save(update_fields=["expires_at"])

    assert api_client_no_csrf.get(_download_url(job.pk)).status_code == 404


@pytest.mark.django_db
def test_download_sets_safe_response_headers(api_client_no_csrf, admin_user, enrollment):
    response = _queue(api_client_no_csrf, admin_user)
    job_id = response.json()["id"]
    download = api_client_no_csrf.get(_download_url(job_id))

    assert download["Content-Disposition"].startswith("attachment;")
    assert download["X-Content-Type-Options"] == "nosniff"
    assert download["Cache-Control"] == "private, no-store"


@pytest.mark.django_db
def test_expires_at_is_set_on_completion(api_client_no_csrf, admin_user, enrollment):
    from apps.reporting.models import EXPORT_RETENTION

    response = _queue(api_client_no_csrf, admin_user)
    job = ExportJob.objects.get(pk=response.json()["id"])

    assert job.expires_at is not None
    assert job.finished_at is not None
    assert job.expires_at - job.finished_at == EXPORT_RETENTION


# ---------------------------------------------------------------------------
# Soft delete
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_soft_delete_hides_a_job_from_the_default_manager(
    api_client_no_csrf, admin_user, enrollment
):
    from apps.common.deletion import soft_delete

    response = _queue(api_client_no_csrf, admin_user)
    job = ExportJob.objects.get(pk=response.json()["id"])

    soft_delete(instance=job, actor=admin_user, reason="No longer needed.")

    assert not ExportJob.objects.filter(pk=job.pk).exists()
    assert ExportJob.all_objects.filter(pk=job.pk).exists()


@pytest.mark.django_db
def test_a_soft_deleted_job_disappears_from_the_list_endpoint(
    api_client_no_csrf, admin_user, enrollment
):
    from apps.common.deletion import soft_delete

    response = _queue(api_client_no_csrf, admin_user)
    job_id = response.json()["id"]
    job = ExportJob.objects.get(pk=job_id)
    soft_delete(instance=job, actor=admin_user, reason="Cleanup.")

    ids = {row["id"] for row in api_client_no_csrf.get(EXPORTS_URL).json()}
    assert job_id not in ids


# ---------------------------------------------------------------------------
# Audit trail
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_every_stage_of_a_completed_job_is_audited(api_client_no_csrf, admin_user, enrollment):
    response = _queue(api_client_no_csrf, admin_user)
    job_id = response.json()["id"]
    api_client_no_csrf.get(_download_url(job_id))

    actions = set(AuditLog.objects.filter(resource_id=job_id).values_list("action", flat=True))
    assert AuditAction.EXPORT_QUEUED in actions
    assert AuditAction.EXPORT_COMPLETED in actions
    assert AuditAction.EXPORT_DOWNLOADED in actions


@pytest.mark.django_db
def test_a_failed_job_is_audited(api_client_no_csrf, admin_user, enrollment, monkeypatch):
    monkeypatch.setattr(writers, "PDF_MAX_ROW_LIMIT", 0)
    response = _queue(api_client_no_csrf, admin_user, fmt=ExportFormat.PDF)
    job_id = response.json()["id"]

    actions = set(AuditLog.objects.filter(resource_id=job_id).values_list("action", flat=True))
    assert AuditAction.EXPORT_FAILED in actions


@pytest.mark.django_db
def test_a_cancelled_job_is_audited(api_client_no_csrf, admin_user, enrollment, monkeypatch):
    monkeypatch.setattr(tasks.run_export, "delay", lambda *a, **k: None)
    response = _queue(api_client_no_csrf, admin_user)
    job_id = response.json()["id"]
    api_client_no_csrf.post(_cancel_url(job_id))

    actions = set(AuditLog.objects.filter(resource_id=job_id).values_list("action", flat=True))
    assert AuditAction.EXPORT_CANCELLED in actions


# ---------------------------------------------------------------------------
# Serialisation: null, not absent, not the string "None"
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_a_queued_jobs_null_fields_serialise_as_json_null(
    api_client_no_csrf, admin_user, enrollment, monkeypatch
):
    monkeypatch.setattr(tasks.run_export, "delay", lambda *a, **k: None)
    response = _queue(api_client_no_csrf, admin_user)
    body = response.json()

    assert body["started_at"] is None
    assert body["finished_at"] is None
    assert body["expires_at"] is None
    assert body["download_url"] is None
    assert body["error"] == ""
    assert body["original_filename"] == ""


@pytest.mark.django_db
def test_requested_by_email_is_null_once_the_account_is_gone(
    api_client_no_csrf, admin_user, unbounded_superadmin, enrollment
):
    """Read back by a platform operator, not by the administrator who queued it.

    A job's centre is its requester's (`access.visible_export_jobs`), so once
    the account is gone the job belongs to no centre and no bounded caller
    reaches it — including the administrator whose account it used to be, whose
    id is exactly the thing that has just been cleared. That is the fail-closed
    reading, and `tests/test_branch_scoping_api.py` states it as a rule of its
    own. What this test is about is unchanged: the serializer renders a missing
    requester as `null` rather than blowing up on the null relation.
    """
    response = _queue(api_client_no_csrf, admin_user)
    job = ExportJob.objects.get(pk=response.json()["id"])

    job.requested_by = None
    job.save(update_fields=["requested_by"])

    api_client_no_csrf.force_login(unbounded_superadmin)
    body = api_client_no_csrf.get(_detail_url(job.pk)).json()
    assert body["requested_by_email"] is None


# ---------------------------------------------------------------------------
# Recorded metadata, and query cost
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_filters_are_recorded_on_the_job(api_client_no_csrf, admin_user, batch, enrollment):
    response = _queue(api_client_no_csrf, admin_user, report_key="attendance", batch=str(batch.pk))
    job = ExportJob.objects.get(pk=response.json()["id"])

    assert job.filters["batch_id"] == str(batch.pk)
    assert job.filters["batch_label"] == batch.code


@pytest.mark.django_db
def test_row_count_checksum_and_size_are_recorded_for_every_completed_job(
    api_client_no_csrf, admin_user, enrollment
):
    response = _queue(api_client_no_csrf, admin_user, fmt=ExportFormat.XLSX)
    job = ExportJob.objects.get(pk=response.json()["id"])

    assert job.row_count >= 1
    assert len(job.checksum) == 64
    assert job.size_bytes > 0
    assert job.original_filename.endswith(".xlsx")


# ---------------------------------------------------------------------------
# Phase 20: the count preflight
# ---------------------------------------------------------------------------


def _count_url(key: str) -> str:
    return f"/api/v1/reports/{key}/count/"


def _export_url(key: str) -> str:
    return f"/api/v1/reports/{key}/export/"


@pytest.mark.django_db
def test_count_matches_the_rows_the_export_actually_produces(
    api_client_no_csrf, admin_user, enrollment
):
    """The one thing this endpoint must never do: disagree with the export
    it is a preflight for. Both numbers come from the same request cycle so
    a flaky, time-dependent report could not make them agree by accident.
    """
    api_client_no_csrf.force_login(admin_user)
    count_response = api_client_no_csrf.get(_count_url("student_progress"))
    assert count_response.status_code == 200
    counted_rows = count_response.json()["rows"]
    assert counted_rows >= 1

    export_response = api_client_no_csrf.get(_export_url("student_progress"))
    assert export_response.status_code == 200
    body = b"".join(export_response.streaming_content).decode("utf-8")
    exported_rows = list(csv.reader(io.StringIO(body)))[1:]  # drop the header
    assert len(exported_rows) == counted_rows


@pytest.mark.django_db
def test_count_is_forbidden_for_a_trainer(api_client_no_csrf, trainer_profile, enrollment):
    api_client_no_csrf.force_login(trainer_profile.user)
    response = api_client_no_csrf.get(_count_url("student_progress"))
    assert response.status_code == 403


@pytest.mark.django_db
def test_count_404s_for_an_unknown_report(api_client_no_csrf, admin_user):
    api_client_no_csrf.force_login(admin_user)
    assert api_client_no_csrf.get(_count_url("not-a-real-report")).status_code == 404


@pytest.mark.django_db
def test_count_respects_the_same_batch_filter_as_the_export(
    api_client_no_csrf, admin_user, batch, enrollment
):
    api_client_no_csrf.force_login(admin_user)
    response = api_client_no_csrf.get(_count_url("attendance"), {"batch": str(batch.pk)})
    assert response.status_code == 200
    assert response.json()["rows"] >= 1


# ---------------------------------------------------------------------------
# Phase 20: the print format
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_print_format_returns_html_not_a_downloadable_file(
    api_client_no_csrf, admin_user, enrollment
):
    api_client_no_csrf.force_login(admin_user)
    response = api_client_no_csrf.get(_export_url("student_progress"), {"as": "print"})
    assert response.status_code == 200
    assert response["Content-Type"].startswith("text/html")
    assert response["Content-Disposition"] == "inline"
    body = response.content.decode("utf-8")
    assert body.startswith("<!doctype html>")
    assert "@media print" in body
    assert "<table>" in body


@pytest.mark.django_db
def test_print_format_still_requires_the_export_capability(
    api_client_no_csrf, trainer_profile, enrollment
):
    api_client_no_csrf.force_login(trainer_profile.user)
    response = api_client_no_csrf.get(_export_url("student_progress"), {"as": "print"})
    assert response.status_code == 403


@pytest.mark.django_db
def test_an_unknown_as_value_is_still_rejected(api_client_no_csrf, admin_user, enrollment):
    api_client_no_csrf.force_login(admin_user)
    response = api_client_no_csrf.get(_export_url("student_progress"), {"as": "doc"})
    assert response.status_code == 400


# ---------------------------------------------------------------------------
# Phase 20: the three new reports, and that each is scoped correctly
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_work_activities_report_only_shows_activities_from_the_callers_branch(
    api_client_no_csrf,
    manager_user,
    unbounded_superadmin,
    enrollment,
    other_branch_enrollment,
):
    """`manager_user` is bounded to `branch` (unlike an admin, who — per the
    owner's 14 September 2026 amendment to D-129, "admin sees every centre" —
    is unbounded and would see both). An activity created against
    `other_branch_enrollment` (a different centre) must never appear in a
    bounded caller's report, even though both activities share the same
    activity type.
    """
    from apps.work.models import ActivityStatus, ActivityType
    from apps.work.services import create_activity

    activity_type = ActivityType.objects.create(
        slug="test-report-type",
        name="Report Test Type",
        category="mentoring",
        allowed_creator_roles=["manager"],
    )
    mine = create_activity(
        actor=manager_user,
        student=enrollment.student,
        activity_type=activity_type,
        enrollment=enrollment,
        title="My branch's activity",
    )
    create_activity(
        actor=unbounded_superadmin,
        student=other_branch_enrollment.student,
        activity_type=activity_type,
        enrollment=other_branch_enrollment,
        title="The other branch's activity",
    )

    api_client_no_csrf.force_login(manager_user)
    response = api_client_no_csrf.get(_export_url("work_activities"))
    assert response.status_code == 200
    body = b"".join(response.streaming_content).decode("utf-8")
    assert enrollment.student.student_id in body
    assert other_branch_enrollment.student.student_id not in body
    assert mine.status == ActivityStatus.DRAFT


@pytest.mark.django_db
def test_deliveries_report_only_shows_deliveries_to_the_callers_branch(
    api_client_no_csrf, manager_user, unbounded_superadmin, student_profile, other_branch_student
):
    """`manager_user` is bounded to `branch` — unlike an admin, who is
    unbounded per the owner's amendment to D-129 and would see both."""
    from apps.communication.models import Delivery, DeliveryState, MessageChannel

    mine = Delivery.objects.create(
        channel=MessageChannel.IN_APP,
        recipient=student_profile.user,
        state=DeliveryState.SENT,
        attempts=1,
    )
    other = Delivery.objects.create(
        channel=MessageChannel.IN_APP,
        recipient=other_branch_student.user,
        state=DeliveryState.SENT,
        attempts=1,
    )

    api_client_no_csrf.force_login(manager_user)
    response = api_client_no_csrf.get(_export_url("deliveries"))
    assert response.status_code == 200
    body = b"".join(response.streaming_content).decode("utf-8")
    assert student_profile.user.get_full_name() in body
    assert other_branch_student.user.get_full_name() not in body
    assert mine.pk != other.pk


@pytest.mark.django_db
def test_deliveries_report_never_leaks_an_otp_send(api_client_no_csrf, admin_user, student_profile):
    """Phase 19's own design (D-051, `apps.communication.services`'s module
    docstring): an OTP send never creates a `Delivery` row at all — it goes
    through `apps.accounts.otp.send_email_code` directly, never through
    `create_deliveries`. So this report needs no extra filtering to keep an
    OTP's variables out; it holds because there is never a row to filter.
    Sending a real code and then checking the report stays honest about
    that, rather than trusting the docstring's claim on its own.
    """
    from apps.accounts.otp import OtpPurpose, send_email_code
    from apps.communication.models import Delivery

    send_email_code(user=student_profile.user, purpose=OtpPurpose.STEP_UP)
    assert not Delivery.objects.filter(recipient=student_profile.user).exists()

    api_client_no_csrf.force_login(admin_user)
    response = api_client_no_csrf.get(_export_url("deliveries"))
    assert response.status_code == 200
    body = b"".join(response.streaming_content).decode("utf-8")
    assert student_profile.user.get_full_name() not in body


@pytest.mark.django_db
def test_deliveries_report_is_forbidden_without_the_communication_capability(
    trainer_profile, student_profile
):
    """The "deliveries" branch of `_queryset_for_user` gates on
    `communication.view_any` itself, in addition to whatever
    `ReportExportView`'s own two outer gates (`report.view_any`,
    `data.export`) already enforce.

    Under the current role model every stock role that clears those two
    outer gates (manager, counsellor, admin, superadmin — D-130 put the
    counsellor beside the manager, "both can do both") also holds
    `communication.view_any`, so there is no real end-to-end account that
    reaches `ReportExportView` and gets refused only at this inner check.
    This calls the domain branch directly instead — a trainer holds neither
    outer capability nor this one, so it still proves the branch's own gate
    independently of whichever roles happen to combine the two today.
    """
    from apps.communication.models import Delivery, DeliveryState, MessageChannel
    from apps.reporting.views import _queryset_for_user

    Delivery.objects.create(
        channel=MessageChannel.IN_APP,
        recipient=student_profile.user,
        state=DeliveryState.SENT,
        attempts=1,
    )

    queryset = _queryset_for_user(trainer_profile.user, "deliveries", None, None)
    assert not queryset.exists()


@pytest.mark.django_db
def test_automation_runs_report_requires_automation_manage(
    api_client_no_csrf, admin_user, counsellor_user
):
    from apps.automation.models import (
        AutomationRule,
        AutomationRuleStatus,
        AutomationRun,
        AutomationRunStatus,
        AutomationTrigger,
    )

    rule = AutomationRule.objects.create(
        name="Test rule",
        trigger=AutomationTrigger.ACTIVITY_COMPLETED,
        status=AutomationRuleStatus.ACTIVE,
    )
    AutomationRun.objects.create(
        rule=rule,
        trigger=AutomationTrigger.ACTIVITY_COMPLETED,
        occurrence_key="test-occurrence-1",
        status=AutomationRunStatus.RAN,
    )

    api_client_no_csrf.force_login(admin_user)
    response = api_client_no_csrf.get(_export_url("automation_runs"))
    assert response.status_code == 200
    body = b"".join(response.streaming_content).decode("utf-8")
    assert "Test rule" in body
    assert "test-occurrence-1" in body

    # A counsellor holds `report.view_any`/`data.export` (D-130) but not
    # `automation.manage` — the report must come back empty for them, the
    # same rule `_queryset_for_user`'s "automation_runs" branch enforces.
    api_client_no_csrf.force_login(counsellor_user)
    response = api_client_no_csrf.get(_export_url("automation_runs"))
    if response.status_code == 200:
        body = b"".join(response.streaming_content).decode("utf-8")
        assert "Test rule" not in body
    else:
        assert response.status_code == 403


@pytest.mark.django_db
def test_automation_runs_report_scrubs_the_error_field(api_client_no_csrf, admin_user):
    from apps.automation.models import (
        AutomationRule,
        AutomationRuleStatus,
        AutomationRun,
        AutomationRunStatus,
        AutomationTrigger,
    )

    rule = AutomationRule.objects.create(
        name="Failing rule",
        trigger=AutomationTrigger.ACTIVITY_COMPLETED,
        status=AutomationRuleStatus.ACTIVE,
    )
    AutomationRun.objects.create(
        rule=rule,
        trigger=AutomationTrigger.ACTIVITY_COMPLETED,
        occurrence_key="test-occurrence-2",
        status=AutomationRunStatus.FAILED,
        error="token=eyJsecretvalue123 while calling the action",
    )

    api_client_no_csrf.force_login(admin_user)
    response = api_client_no_csrf.get(_export_url("automation_runs"))
    body = b"".join(response.streaming_content).decode("utf-8")
    assert "eyJsecretvalue123" not in body
    assert "token=" in body


@pytest.mark.django_db
def test_the_list_endpoint_has_a_bounded_query_count(
    api_client_no_csrf, admin_user, enrollment, django_assert_max_num_queries
):
    for _ in range(5):
        ExportJob.objects.create(
            report_key="student_progress",
            format=ExportFormat.CSV,
            requested_by=admin_user,
            status=ExportStatus.COMPLETED,
        )
    api_client_no_csrf.force_login(admin_user)
    with django_assert_max_num_queries(10):
        response = api_client_no_csrf.get(EXPORTS_URL)
    assert response.status_code == 200
    assert len(response.json()) >= 5
